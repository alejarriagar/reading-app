import io
import json

import pytest
from openpyxl import Workbook

from backend import accounts, store, xlsx_import
from backend import atlas as atlas_module
from backend.atlas import Atlas


def example_file():
    workbook = Workbook()
    books = workbook.active
    books.title = "Libros"
    books.append([
        "Autor", "Título", "Género", "Volumenes", "Lectura",
        "Valoración", "Frases", "Comentario", "Precio", "Emplazamiento",
    ])
    books.append(["ABÉCASSIS, Eliette", "Qumrán", "Novela", 1, "Sí", 8.5, "Frase privada", "Comentario privado", 1500, "Calpe"])
    books.append(["Academia", "Cocina", "Cocina", 1, "Parcial", None, None, None, 900, "Estanterías cocina"])
    books.append(["2E", "Cacúmen. Revista Lúdica", "Lógica", 2, "Sí", None, None, None, None, "Obra"])

    magazines = workbook.create_sheet("Revistas")
    magazines.append(["Título", "Número", "Género", "Volumenes", "Lectura", "Valoración", "Frases", "Comentario", "Precio", "Emplazamiento"])
    magazines.append(["Revista de prueba", "12", "Ciencia", 1, "Sí", 9, None, "Nota privada revista", 500, "Obra"])

    old = workbook.create_sheet("Hoja3")
    old.append(["Título", "Autor", "Género", "Volumenes", "Lectura", "Valoración", "Frases", "Comentario", "Precio", "Emplazamiento"])
    old.append(["Libro antiguo", "Autor antiguo", "Novela", 1, "No", None, None, None, 1000, "Biblioteca"])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture
def lab(tmp_path, monkeypatch):
    monkeypatch.setattr(accounts, "ACCOUNTS_FILE", tmp_path / "accounts.db")
    monkeypatch.setattr(accounts, "OWNER_DIR", tmp_path)
    monkeypatch.setattr(accounts, "PROFILES_DIR", tmp_path / "profiles")
    monkeypatch.setattr(store, "DB_FILE", tmp_path / "atlas.db")
    monkeypatch.setattr(store, "LEGACY_FILE", tmp_path / "legacy.json")
    monkeypatch.setattr(store, "BACKUP_DIR", tmp_path / "backups")
    monkeypatch.setattr(xlsx_import, "ROOT", tmp_path)

    catalog = {"regions": {}, "authors": {}, "books": []}
    (tmp_path / "catalog.json").write_text(json.dumps(catalog), encoding="utf-8")
    (tmp_path / "paths.json").write_text('{"paths":[]}', encoding="utf-8")

    monkeypatch.setattr(atlas_module, "CATALOG_FILE", tmp_path / "catalog.json")
    monkeypatch.setattr(atlas_module, "PATHS_FILE", tmp_path / "paths.json")
    monkeypatch.setattr(atlas_module, "WIKIDATA_FILE", tmp_path / "missing.json")

    owner = accounts.create_user("Yo", "clave123", is_owner=True)
    dad = accounts.create_user("Papá", "clave456")

    def open_profile(user):
        db, backups, legacy = accounts.user_files(user)
        store.DB_FILE, store.BACKUP_DIR, store.LEGACY_FILE = db, backups, legacy
        if "user_id" in __import__("inspect").signature(Atlas).parameters:
            result = Atlas(user_id=user["id"])
        else:
            result = Atlas()
            result.user_id = user["id"]
        return result

    return owner, dad, open_profile


def test_parse_three_sheets():
    parsed = xlsx_import.parse(example_file())
    assert len(parsed["rows"]) == 5
    first = parsed["rows"][0]
    assert first["author"] == "Eliette Abécassis"
    assert first["house"] == "Calpe"
    assert parsed["rows"][2]["mediaKind"] == "MAGAZINE"


def test_import_profile_inventory_and_repeat(lab):
    owner, dad, open_profile = lab

    preview = xlsx_import.preview(example_file(), owner, dad["id"])
    result = xlsx_import.commit(preview["rows"], owner, dad["id"], open_profile)
    assert result["created"] == 5

    dad_atlas = open_profile(dad)
    snapshot = dad_atlas.snapshot()

    assert snapshot["summary"]["booksRead"] == 1
    assert snapshot["summary"]["reading"] == 0
    assert snapshot["summary"]["byYear"] == {}
    assert len(dad_atlas.profile["magazine_records"]) == 2

    book = next(b for b in snapshot["books"] if b["title"] == "Qumrán")
    assert book["rating"] == 8.5
    assert book["readings"][0]["finished"] is None
    assert "Comentario privado" in book["notes"]

    from backend import library
    copies = library.list_copies()
    assert len(copies) == 5
    assert all("quotes" not in c["excel"] and "comment" not in c["excel"] for c in copies)
    assert any(c["excel"]["price"] == "1500" for c in copies)

    assert open_profile(owner).snapshot()["summary"]["booksRead"] == 0

    repeated = xlsx_import.preview(example_file(), owner, dad["id"])
    assert all(row["alreadyImported"] for row in repeated["rows"])

    again = xlsx_import.commit(repeated["rows"], owner, dad["id"], open_profile)
    assert again["created"] == 0
    assert again["personalRows"] == 0
    assert len(library.list_copies()) == 5


def test_import_permissions(lab):
    owner, dad, _ = lab
    with pytest.raises(PermissionError):
        xlsx_import.preview(example_file(), dad, owner["id"])



def test_batch_rebuilds_catalog_once(lab, monkeypatch):
    owner, dad, open_profile = lab
    atlas = open_profile(dad)
    rows = xlsx_import.parse(example_file())["rows"]
    pairs = [(f"copy-{index}", row) for index, row in enumerate(rows)]

    calls = []
    original_index = atlas._index

    def counted_index():
        calls.append(1)
        return original_index()

    monkeypatch.setattr(atlas, "_index", counted_index)

    progress = []
    processed, warnings = xlsx_import.merge_personal(
        atlas, pairs, progress=progress.append,
    )

    assert processed == 5
    assert len(calls) == 1
    assert atlas.snapshot()["summary"]["booksRead"] == 1
    assert atlas.snapshot()["summary"]["reading"] == 0
    assert atlas.snapshot()["summary"]["byYear"] == {}
    assert progress[-1]["stage"] == "save"
    assert progress[-1]["done"] == 1


def test_batch_reuses_new_author_and_work(lab):
    owner, dad, open_profile = lab
    atlas = open_profile(dad)

    first = xlsx_import.parse(example_file())["rows"][0]
    second = dict(first)
    second["sourceRow"] = 200
    second["title"] = "Otra obra"
    second["author"] = first["author"].upper()

    repeated_work = dict(first)
    repeated_work["sourceRow"] = 201

    processed, warnings = xlsx_import.merge_personal(
        atlas,
        [("c1", first), ("c2", second), ("c3", repeated_work)],
    )

    assert processed == 3
    assert len(atlas.profile["custom_authors"]) == 1
    assert len(atlas.profile["custom_books"]) == 2
    assert len([
        reading for reading in atlas.profile["readings"]
        if reading["status"] == "done"
    ]) == 2


def test_job_status_permissions(lab, monkeypatch):
    import time
    from backend import xlsx_jobs

    owner, dad, _ = lab
    monkeypatch.setattr(xlsx_jobs, "_JOBS", {})

    def task(progress):
        progress({
            "stage": "save",
            "label": "Prueba",
            "done": 1,
            "total": 1,
            "percent": 99,
        })
        return {"ok": True}

    job_id = xlsx_jobs.start(owner, dad["id"], task)["jobId"]

    deadline = time.time() + 3
    while time.time() < deadline:
        status = xlsx_jobs.status(job_id, owner)
        if status["state"] == "done":
            break
        time.sleep(0.01)

    assert status["state"] == "done"
    assert status["progress"]["percent"] == 100
    assert xlsx_jobs.status(job_id, dad)["result"]["ok"]

    with pytest.raises(PermissionError):
        xlsx_jobs.status(job_id, {"id": "otro", "isOwner": False})

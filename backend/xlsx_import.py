import copy
import hashlib
import io
import json
import re
import shutil
import sqlite3
import uuid
import zipfile
from contextlib import closing
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from backend import accounts, excel_meta, library, store
from backend.atlas import DEFAULT_PROFILE, norm

ROOT = Path(__file__).resolve().parents[1]
MAX_BYTES = 8 * 1024 * 1024
MAX_EXPANDED = 40 * 1024 * 1024
MAX_ROWS = 3000

HEADERS = {
    "authorOriginal": {"autor", "autora"},
    "title": {"titulo"},
    "genre": {"genero"},
    "volumes": {"volumenes", "volumen"},
    "reading": {"lectura"},
    "ratingOriginal": {"valoracion"},
    "quotes": {"frases"},
    "comment": {"comentario"},
    "price": {"precio"},
    "location": {"emplazamiento"},
    "number": {"numero"},
}

GENRES = {
    "novela": "NOVELA",
    "relato": "RELATO",
    "relatos": "RELATO",
    "cuento": "RELATO",
    "cuentos": "RELATO",
    "poesia": "POESIA",
    "teatro": "TEATRO",
    "ensayo": "ENSAYO",
}

ROW_FIELDS = (
    "title", "author", "authorOriginal", "genre", "volumes", "reading",
    "ratingOriginal", "quotes", "comment", "price", "location", "number",
    "sourceSheet", "sourceRow", "mediaKind", "house",
)


def text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    return str(value).strip()


def author_name(value):
    value = text(value)
    if value.count(",") == 1:
        surname, first = [part.strip() for part in value.split(",", 1)]
        if surname and first and len(first.split()) <= 5:
            surname = surname.title() if surname.isupper() else surname
            return f"{first} {surname}"
    return value


def reading_state(value):
    key = norm(value)
    if key in {"si", "yes", "leido", "leida", "true"}:
        return "READ"
    if key in {"parcial", "parcialmente"}:
        return "PARTIAL"
    if key in {"", "no", "false", "sin leer", "no leido"}:
        return "UNREAD"
    return "UNKNOWN"


def rating(value):
    value = text(value)
    if not value:
        return None
    try:
        number = float(value.replace(",", "."))
    except ValueError:
        return None
    return number if 0 <= number <= 10 else None


def row_key(row):
    # El registro guarda el hash, no las frases o comentarios personales.
    original = {
        key: row.get(key, "")
        for key in ROW_FIELDS
        if key not in {"mediaKind", "house"}
    }
    return hashlib.sha256(
        json.dumps(original, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()


def clean_row(raw):
    if not isinstance(raw, dict):
        raise ValueError("Fila no válida.")

    row = {key: text(raw.get(key)) for key in ROW_FIELDS}
    row["sourceRow"] = int(raw.get("sourceRow") or 0)

    if row["sourceSheet"] not in {"Libros", "Revistas", "Hoja3"}:
        raise ValueError("Pestaña no reconocida.")
    if row["sourceRow"] < 2:
        raise ValueError("Número de fila no válido.")
    if not row["title"]:
        raise ValueError("Falta el título.")
    if len(row["title"]) > 300 or len(row["author"]) > 200:
        raise ValueError("Título o autor demasiado largo.")
    if row["mediaKind"] not in {"BOOK", "MAGAZINE"}:
        raise ValueError("Tipo de registro no válido.")
    if row["house"] not in {"Pozuelo", "Calpe"}:
        raise ValueError("La casa debe ser Pozuelo o Calpe.")
    if any(len(row[key]) > 20000 for key in ("quotes", "comment")):
        raise ValueError("Frases o comentario demasiado largos.")

    row["allowExtra"] = bool(raw.get("allowExtra"))
    return row


def parse(data):
    if not data or len(data) > MAX_BYTES:
        raise ValueError("El archivo está vacío o supera 8 MB.")

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED:
                raise ValueError("El contenido del Excel es demasiado grande.")
            if any(item.filename.lower().endswith("vbaproject.bin")
                   for item in archive.infolist()):
                raise ValueError("Usa un .xlsx sin macros.")
    except zipfile.BadZipFile as exc:
        raise ValueError("El archivo no es un .xlsx válido.") from exc

    try:
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=True, keep_links=False,
        )
    except Exception as exc:
        raise ValueError("No se ha podido leer el Excel.") from exc

    records, warnings, sheets = [], [], []

    try:
        for sheet_name in ("Libros", "Revistas", "Hoja3"):
            if sheet_name not in workbook.sheetnames:
                warnings.append(f"No existe la pestaña «{sheet_name}».")
                continue

            sheet = workbook[sheet_name]
            if sheet.max_row > 10000:
                raise ValueError(
                    f"«{sheet_name}» tiene más de 10.000 filas, posiblemente por formato "
                    "sobrante. Guarda una copia solo con las filas utilizadas."
                )

            values = sheet.iter_rows(
                max_row=sheet.max_row, max_col=min(sheet.max_column, 40),
                values_only=True,
            )
            header = next(values, None)
            if not header:
                continue

            normalized = [norm(text(value)) for value in header]
            mapping = {}

            for field, aliases in HEADERS.items():
                for index, name in enumerate(normalized):
                    if name in aliases:
                        mapping[field] = index
                        break

            if "title" not in mapping:
                raise ValueError(
                    f"No encuentro la columna Título en «{sheet_name}»."
                )

            count = 0
            for row_number, values_row in enumerate(values, start=2):
                if not any(value is not None and text(value) for value in values_row):
                    continue

                raw = {
                    field: text(values_row[index]) if index < len(values_row) else ""
                    for field, index in mapping.items()
                }
                if not raw.get("title"):
                    warnings.append(f"{sheet_name}, fila {row_number}: sin título; omitida.")
                    continue

                raw["sourceSheet"] = sheet_name
                raw["sourceRow"] = row_number
                raw["author"] = author_name(raw.get("authorOriginal"))

                # Es una propuesta visible y modificable en la vista previa.
                raw["mediaKind"] = (
                    "MAGAZINE"
                    if sheet_name == "Revistas"
                    or re.search(r"\brevista\b", norm(raw["title"]))
                    else "BOOK"
                )

                raw["house"] = (
                    "Calpe"
                    if re.search(r"\bcalpe\b", norm(raw.get("location", "")))
                    else "Pozuelo"
                )

                record = clean_row(raw)
                records.append(record)
                count += 1

                if len(records) > MAX_ROWS:
                    raise ValueError("Máximo 3.000 registros por importación.")

            sheets.append({"name": sheet_name, "rows": count})

    finally:
        workbook.close()

    if not records:
        raise ValueError("No se han encontrado registros con título.")

    return {"rows": records, "warnings": warnings, "sheets": sheets}


def authorized_target(actor, user_id):
    target = next(
        (user for user in accounts.list_users() if user["id"] == user_id), None
    )
    if target is None:
        raise ValueError("El perfil elegido no existe.")
    if actor["id"] != target["id"] and not actor.get("isOwner"):
        raise PermissionError("Solo el perfil principal puede importar para otro perfil.")
    return target


def peek_profile(user):
    path, _, _ = accounts.user_files(user)
    if not path.exists():
        return copy.deepcopy(DEFAULT_PROFILE)

    with closing(sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True)) as conn:
        return store._read(conn, DEFAULT_PROFILE)


def matching_copies(row, copies, house_names):
    result = []

    for item in copies:
        if house_names.get(item["houseId"]) != row["house"]:
            continue
        if norm(item["title"]) != norm(row["title"]):
            continue
        if item.get("mediaKind", "BOOK") != row["mediaKind"]:
            continue

        if row["mediaKind"] == "MAGAZINE":
            existing_number = (item.get("excel") or {}).get("number", "")
            if norm(existing_number) != norm(row["number"]):
                continue
        elif norm(author_name(item.get("author"))) != norm(row["author"]):
            continue

        result.append(item)

    return result


def preview(data, actor, user_id):
    target = authorized_target(actor, user_id)
    parsed = parse(data)
    copies = library.list_copies()
    houses = {house["id"]: house["name"] for house in library.list_houses()}
    existing_ids = {item["id"] for item in copies}
    profile = peek_profile(target)
    completed = set(profile.get("_xlsx_rows", []))

    with closing(excel_meta.connect()) as conn:
        ledger = dict(conn.execute(
            "SELECT row_key, copy_id FROM excel_import_row WHERE user_id = ?",
            (target["id"],),
        ))

    seen = set()

    for row in parsed["rows"]:
        key = row_key(row)
        repeated = key in ledger and ledger[key] in existing_ids and key in completed
        matches = matching_copies(row, copies, houses)
        issues = []

        identity = (
            row["house"], row["mediaKind"], norm(row["title"]),
            norm(row["author"]), norm(row["number"]),
        )

        if identity in seen:
            issues.append("Otra fila del archivo tiene la misma identidad; revisa si es otro ejemplar.")
        seen.add(identity)

        if row["ratingOriginal"] and rating(row["ratingOriginal"]) is None:
            issues.append("Valoración fuera de 0–10 o no numérica: se conserva como texto.")

        if reading_state(row["reading"]) == "UNKNOWN":
            issues.append("Lectura no reconocida: se conserva sin marcar como leído.")

        if len(matches) > 1:
            issues.append("Hay varios ejemplares coincidentes; elige crear otro o deja esta fila sin seleccionar.")

        row["alreadyImported"] = repeated
        row["existingCopies"] = len(matches)
        row["issues"] = issues
        row["selected"] = not repeated and len(matches) <= 1

    return {
        **parsed,
        "target": {"id": target["id"], "name": target["name"]},
    }


def backup_before_import(target):
    directory = ROOT / "backups" / (
        f"excel-{datetime.now():%Y%m%d-%H%M%S}-{uuid.uuid4().hex[:6]}"
    )
    directory.mkdir(parents=True, exist_ok=True)

    profile_path, _, _ = accounts.user_files(target)
    paths = (
        (accounts.ACCOUNTS_FILE, "accounts.db"),
        (profile_path, "profile.db"),
    )

    for source, name in paths:
        if not source.exists():
            continue
        with closing(sqlite3.connect(source.resolve().as_uri() + "?mode=ro", uri=True)) as src:
            with closing(sqlite3.connect(directory / name)) as dst:
                src.backup(dst)

    (directory / "target.json").write_text(
        json.dumps({"id": target["id"], "name": target["name"]}, ensure_ascii=False),
        encoding="utf-8",
    )
    return directory.name


def physical_source(row):
    # Las notas de lectura NO se guardan en la biblioteca compartida.
    return {
        key: row[key]
        for key in (
            "sourceSheet", "sourceRow", "mediaKind", "genre", "volumes",
            "price", "location", "number", "authorOriginal",
        )
    } | {"currency": "ESP"}


def personal_source(row):
    return {
        key: row[key]
        for key in (
            "sourceSheet", "sourceRow", "reading", "ratingOriginal",
            "quotes", "comment", "genre", "authorOriginal",
        )
    }


def find_work(atlas, row):
    title = norm(row["title"])
    author = norm(row["author"] or "Autor desconocido")
    matches = [
        item for item in atlas._library
        if title in item.get("title_keys", {item["title_key"]})
        and author in item.get("author_keys", {item["author_key"]})
    ]
    ids = list(dict.fromkeys(item["id"] for item in matches))
    return ids


def _xlsx_progress(callback, stage, done, total):
    if callback is None:
        return

    labels = {
        "prepare": "Preparando y guardando copia de seguridad",
        "inventory": "Revisando el inventario",
        "profile": "Preparando los datos personales",
        "save": "Guardando el perfil",
    }
    bands = {
        "prepare": (0, 5),
        "inventory": (5, 50),
        "profile": (55, 40),
        "save": (95, 4),
    }

    base, width = bands[stage]
    fraction = min(1, max(0, done / total)) if total else 1

    callback({
        "stage": stage,
        "label": labels[stage],
        "done": done,
        "total": total,
        "percent": min(99, round(base + width * fraction)),
    })


def _xlsx_work_index(atlas):
    index = {}

    for item in atlas._library:
        titles = item.get("title_keys", {item["title_key"]})
        names = item.get("author_keys", {item["author_key"]})

        for title in titles:
            for name in names:
                ids = index.setdefault((title, name), [])
                if item["id"] not in ids:
                    ids.append(item["id"])

    return index


def _xlsx_author_index(atlas):
    index = {}

    for author_id, author in atlas.authors.items():
        names = author.get("_catalog_names", [author["name"]])
        for name in names:
            index.setdefault(norm(name), author_id)

    return index


def _xlsx_add_work(profile, row, work_index, author_index):
    """Crear registros privados sin reconstruir el catálogo en cada fila."""
    title = row["title"].strip()
    author_name = (row["author"] or "Autor desconocido").strip()
    author_key = norm(author_name)

    author_id = author_index.get(author_key)

    if author_id is None:
        author_id = f"custom-author-{uuid.uuid4().hex[:12]}"
        profile["custom_authors"].append({
            "id": author_id,
            "name": author_name,
            "country": "No especificado",
            "period": "No especificado",
            "movements": [],
            "tier": "SPECIALIZED",
            "regions": [],
            "related": [],
        })
        author_index[author_key] = author_id
    else:
        profile["author_region_additions"].setdefault(author_id, [])

    book_id = f"custom-book-{uuid.uuid4().hex[:12]}"

    profile["custom_books"].append({
        "id": book_id,
        "title": title,
        "author": author_id,
        "regions": [],
    })
    profile["book_statuses"][book_id] = "LOCKED"
    work_index.setdefault((norm(title), author_key), []).append(book_id)

    return book_id


def merge_personal(atlas, pairs, progress=None):
    profile = atlas.profile
    completed = set(profile.get("_xlsx_rows", []))
    warnings = []
    processed = 0

    work_index = _xlsx_work_index(atlas)
    author_index = _xlsx_author_index(atlas)
    total = len(pairs)

    _xlsx_progress(progress, "profile", 0, total)

    for position, (copy_id, row) in enumerate(pairs, start=1):
        key = row_key(row)

        if key in completed:
            _xlsx_progress(progress, "profile", position, total)
            continue

        source = personal_source(row)
        state = reading_state(row["reading"])
        number = rating(row["ratingOriginal"])

        if row["mediaKind"] == "MAGAZINE":
            profile.setdefault("magazine_records", {})[copy_id] = source
            completed.add(key)
            processed += 1
            _xlsx_progress(progress, "profile", position, total)
            continue

        ids = work_index.get((
            norm(row["title"]),
            norm(row["author"] or "Autor desconocido"),
        ), [])

        if len(ids) > 1:
            warnings.append(
                f"«{row['title']}»: varios registros del catálogo; "
                "el ejemplar se ha guardado, pero su lectura queda pendiente de vincular."
            )
            _xlsx_progress(progress, "profile", position, total)
            continue

        book_id = ids[0] if ids else _xlsx_add_work(
            profile, row, work_index, author_index,
        )

        metadata = dict(profile["book_metadata"].get(book_id, {}))
        metadata.setdefault("year", None)
        metadata.setdefault("rating", None)
        metadata.setdefault("notes", "")

        old_state = profile["book_statuses"].get(book_id, "LOCKED")

        if state == "READ":
            profile["book_statuses"][book_id] = "READ"

            done = [
                reading for reading in profile.get("readings", [])
                if reading["book"] == book_id and reading["status"] == "done"
            ]

            if not done:
                opened = [
                    reading for reading in profile.get("readings", [])
                    if reading["book"] == book_id and reading["status"] == "open"
                ]
                if opened:
                    warnings.append(
                        f"«{row['title']}»: se conserva su lectura en curso "
                        "y se registra la lectura anterior sin fecha."
                    )

                profile.setdefault("readings", []).append({
                    "id": uuid.uuid4().hex[:12],
                    "book": book_id,
                    "status": "done",
                    "started": None,
                    "finished": None,
                    "reread": False,
                    "note": "Lectura importada de Excel sin fecha conocida.",
                })

            if number is not None:
                if metadata.get("rating") is None:
                    metadata["rating"] = number
                elif metadata["rating"] != number:
                    warnings.append(
                        f"«{row['title']}»: se conserva su puntuación anterior; "
                        "la del Excel queda en los datos originales."
                    )

        elif state == "PARTIAL" and old_state == "LOCKED":
            profile["book_statuses"][book_id] = "DISCOVERED"

        if not metadata.get("genre") and norm(row["genre"]) in GENRES:
            metadata["genre"] = GENRES[norm(row["genre"])]

        blocks = []
        if state == "PARTIAL":
            blocks.append("Lectura parcial según el Excel.")
        if row["quotes"]:
            blocks.append("Frases:\n" + row["quotes"])
        if row["comment"]:
            blocks.append("Comentario:\n" + row["comment"])

        if blocks:
            additional = "\n\n".join(blocks)
            previous = metadata.get("notes") or ""
            combined = (previous + "\n\n" + additional).strip()

            if len(combined) <= 5000:
                metadata["notes"] = combined
            else:
                warnings.append(
                    f"«{row['title']}»: notas extensas conservadas completas "
                    "en los datos originales del Excel."
                )

        metadata.setdefault("excelSources", []).append(source)
        profile["book_metadata"][book_id] = metadata
        completed.add(key)
        processed += 1

        _xlsx_progress(progress, "profile", position, total)

    profile["_xlsx_rows"] = sorted(completed)

    _xlsx_progress(progress, "save", 0, 1)

    if processed:
        atlas.label = "Importar lecturas y datos personales desde Excel"
        # Un guardado y una reconstrucción del catálogo al final.
        atlas.save()

    _xlsx_progress(progress, "save", 1, 1)

    return processed, warnings


def commit(rows, actor, user_id, open_atlas, progress=None):
    target = authorized_target(actor, user_id)

    if not rows or len(rows) > MAX_ROWS:
        raise ValueError("Selecciona entre 1 y 3.000 filas.")

    cleaned = [clean_row(row) for row in rows]
    _xlsx_progress(progress, "prepare", 0, 1)
    backup_name = backup_before_import(target)
    _xlsx_progress(progress, "prepare", 1, 1)

    # Primero preparar el inventario; después aplicar el perfil.
    # El registro por fila permite repetir un lote interrumpido.
    atlas = open_atlas(target)
    copies = library.list_copies()
    house_names = {h["id"]: h["name"] for h in library.list_houses()}
    users = {u["id"] for u in accounts.list_users()}
    created, reused, skipped = 0, 0, 0
    pairs, warnings = [], []
    _xlsx_progress(progress, "inventory", 0, len(cleaned))

    with closing(excel_meta.connect()) as conn:
        with conn:
            house_ids = {name: identifier for identifier, name in house_names.items()}

            for row in cleaned:
                if row["house"] not in house_ids:
                    identifier = library._insert_house(conn, row["house"])
                    house_ids[row["house"]] = identifier
                    house_names[identifier] = row["house"]

            ledger = dict(conn.execute(
                "SELECT row_key, copy_id FROM excel_import_row WHERE user_id = ?",
                (target["id"],),
            ))

            for position, row in enumerate(cleaned, start=1):
                key = row_key(row)
                known = ledger.get(key)
                copy_id = known if any(c["id"] == known for c in copies) else None

                if copy_id:
                    skipped += 1
                else:
                    matches = matching_copies(row, copies, house_names)

                    if len(matches) > 1 and not row["allowExtra"]:
                        warnings.append(
                            f"«{row['title']}»: varios ejemplares coincidentes; fila omitida."
                        )
                        _xlsx_progress(progress, "inventory", position, len(cleaned))
                        continue

                    if matches and not row["allowExtra"]:
                        copy_id = matches[0]["id"]
                        reused += 1
                    else:
                        ids = find_work(atlas, row) if row["mediaKind"] == "BOOK" else []
                        cover = None
                        if len(ids) == 1:
                            book = atlas.books_by_id[ids[0]]
                            name = atlas.authors.get(book["author"], {}).get("name", "")
                            cover = atlas.cover_for(book["title"], name)

                        payload = library.clean_copy({
                            "title": row["title"],
                            "author": row["author"],
                            "house_id": house_ids[row["house"]],
                            "owner_id": target["id"],
                            "cover": cover or "",
                        }, set(house_names), users)

                        copy_id = library._insert_copy(conn, payload, actor["id"])
                        copies.append({
                            "id": copy_id,
                            "title": row["title"],
                            "author": row["author"],
                            "houseId": house_ids[row["house"]],
                            "mediaKind": row["mediaKind"],
                            "excel": physical_source(row),
                        })
                        created += 1

                    conn.execute(
                        "INSERT INTO excel_import_row (user_id, row_key, copy_id) "
                        "VALUES (?, ?, ?) ON CONFLICT(user_id, row_key) DO UPDATE SET "
                        "copy_id = excluded.copy_id",
                        (target["id"], key, copy_id),
                    )
                    ledger[key] = copy_id

                conn.execute(
                    "INSERT OR IGNORE INTO excel_copy_source (copy_id, row_key, data) "
                    "VALUES (?, ?, ?)",
                    (
                        copy_id, key,
                        json.dumps(physical_source(row), ensure_ascii=False),
                    ),
                )
                pairs.append((copy_id, row))
                _xlsx_progress(progress, "inventory", position, len(cleaned))

    processed, personal_warnings = merge_personal(atlas, pairs, progress)
    warnings.extend(personal_warnings)

    return {
        "created": created,
        "reused": reused,
        "alreadyRecorded": skipped,
        "personalRows": processed,
        "warnings": warnings,
        "backup": backup_name,
        "targetName": target["name"],
    }
